"""Load ground-truth sheets from the working workbook.

This module is the only place that may open those sheets.
"""

from __future__ import annotations

import os
from pathlib import Path

from openpyxl import load_workbook

DEFAULT_DATASET_ROOT = Path(
    "/Users/yuhan/Downloads/Ylookup Hackathon Datasets/"
    "01-bank-statements-to-journal-entries"
)


def workbook_path() -> Path | None:
    raw = Path("data/raw")
    local = next(raw.glob("*.xlsx"), None) if raw.is_dir() else None
    if local:
        return local
    root = Path(os.environ.get("DATASET_ROOT", DEFAULT_DATASET_ROOT))
    matches = list((root / "workbook").glob("*.xlsx"))
    return matches[0] if matches else None


def _sheet_rows(path: Path, wanted: str) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        actual = next((name for name in workbook.sheetnames if name.strip() == wanted), None)
        if actual is None:
            return []
        rows = workbook[actual].iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            return []
        fields = [str(cell).strip() if cell is not None else f"c{i}" for i, cell in enumerate(header)]
        records: list[dict[str, str]] = []
        for raw in rows:
            record = {
                fields[i]: "" if raw[i] is None else str(raw[i]).strip()
                for i in range(len(fields))
            }
            if any(record.values()):
                records.append(record)
        return records
    finally:
        workbook.close()


def load_ground_truth() -> dict[str, list[dict[str, str]]]:
    path = workbook_path()
    if path is None:
        return {"staging": [], "journals": []}
    return {
        "staging": _sheet_rows(path, "Staging Sheet"),
        "journals": _sheet_rows(path, "DIU"),
    }
