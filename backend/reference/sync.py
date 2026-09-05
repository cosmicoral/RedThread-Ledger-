"""Copy hackathon PDFs and extract allowlisted reference sheets to CSV."""

from __future__ import annotations

import csv
import shutil
from pathlib import Path

from openpyxl import load_workbook

from reference.constants import ALLOWED_SHEETS
from settings import settings

DEFAULT_DATASET_ROOT = Path(
    "/Users/yuhan/Downloads/Ylookup Hackathon Datasets/"
    "01-bank-statements-to-journal-entries"
)


def _dataset_root() -> Path:
    import os

    return Path(os.environ.get("DATASET_ROOT", DEFAULT_DATASET_ROOT))


def sync_statements(dataset_root: Path, dest: Path) -> int:
    source = dataset_root / "statements"
    dest.mkdir(parents=True, exist_ok=True)
    copied = 0
    if not source.is_dir():
        return copied
    for pdf in sorted(source.glob("*.pdf")):
        shutil.copy2(pdf, dest / pdf.name)
        copied += 1
    return copied


def _workbook_path(dataset_root: Path) -> Path | None:
    workbook_dir = dataset_root / "workbook"
    if not workbook_dir.is_dir():
        return None
    matches = list(workbook_dir.glob("*.xlsx"))
    return matches[0] if matches else None


def extract_reference_csvs(workbook_path: Path, dest: Path) -> list[str]:
    dest.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    written: list[str] = []
    try:
        present = {name.strip(): name for name in workbook.sheetnames}
        for sheet_name, filename in ALLOWED_SHEETS.items():
            actual = present.get(sheet_name)
            if actual is None:
                continue
            rows = workbook[actual].iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                continue
            fieldnames = [str(cell).strip() if cell is not None else "" for cell in header]
            out_path = dest / filename
            with out_path.open("w", newline="", encoding="utf-8") as handle:
                writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
                writer.writeheader()
                for raw in rows:
                    record = {
                        fieldnames[i]: "" if raw[i] is None else str(raw[i]).strip()
                        for i in range(len(fieldnames))
                    }
                    if any(record.values()):
                        writer.writerow(record)
            written.append(sheet_name)
    finally:
        workbook.close()
    return written


def sync(dataset_root: Path | None = None) -> dict[str, object]:
    root = dataset_root or _dataset_root()
    statements_dir = Path(settings.hackathon_statements_dir)
    reference_dir = Path(settings.hackathon_reference_dir)
    raw_dir = Path(settings.data_dir) / "raw"

    pdf_count = sync_statements(root, statements_dir)
    workbook_path = _workbook_path(root)
    sheets: list[str] = []
    if workbook_path is not None:
        raw_dir.mkdir(parents=True, exist_ok=True)
        shutil.copy2(workbook_path, raw_dir / workbook_path.name)
        sheets = extract_reference_csvs(workbook_path, reference_dir)

    return {
        "dataset_root": str(root),
        "pdfs": pdf_count,
        "sheets": sheets,
    }


if __name__ == "__main__":
    result = sync()
    print(f"Synced {result['pdfs']} PDFs from {result['dataset_root']}")
    print(f"Extracted {len(result['sheets'])} reference sheets: {', '.join(result['sheets'])}")
